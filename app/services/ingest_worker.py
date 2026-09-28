"""
Worker ULTRA LIGERO - procesa documentos de a poco.
Soporta PDF, DOCX, XLSX y TXT.
Usa requests directo a OpenAI y encripta antes de guardar.
"""
import sys
import os
import json
import gc
import base64
import hashlib

# Setup path
from pathlib import Path
backend_dir = Path(__file__).parent.parent.parent
sys.path.insert(0, str(backend_dir))

from dotenv import load_dotenv
load_dotenv(backend_dir / ".env")

import requests
from supabase import create_client


# Encriptación inline
def encrypt_content(text: str) -> str:
    """Encripta el contenido si ENCRYPTION_KEY está configurada"""
    encryption_key = os.getenv("ENCRYPTION_KEY")
    if not encryption_key or not text:
        return text

    try:
        from cryptography.fernet import Fernet
        key_bytes = hashlib.sha256(encryption_key.encode()).digest()
        fernet_key = base64.urlsafe_b64encode(key_bytes)
        fernet = Fernet(fernet_key)
        encrypted = fernet.encrypt(text.encode('utf-8'))
        return "ENC:" + encrypted.decode('utf-8')
    except Exception as e:
        print(f"Warning: No se pudo encriptar: {e}", file=sys.stderr)
        return text


# Config conservador
MAX_PAGES = 50
MAX_TEXT_PER_PAGE = 8000
CHUNK_SIZE = 800


def get_openai_embedding(text: str, api_key: str) -> list:
    """Llama a OpenAI directamente con requests (sin SDK)"""
    response = requests.post(
        "https://api.openai.com/v1/embeddings",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json"
        },
        json={
            "model": "text-embedding-3-small",
            "input": text
        },
        timeout=30
    )

    if response.status_code != 200:
        raise Exception(f"OpenAI error: {response.text}")

    data = response.json()
    return data["data"][0]["embedding"]


def split_text(text: str) -> list:
    """Divide texto en chunks pequeños"""
    if len(text) <= CHUNK_SIZE:
        return [text] if text.strip() else []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))

        # Cortar en espacio si es posible
        if end < len(text):
            space = text.rfind(' ', start, end)
            if space > start:
                end = space

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        start = end

    return chunks


# --- Extractores por formato ---

def extract_pages_pdf(file_path: str) -> list:
    """Extrae páginas de un PDF. Retorna [{page_number, text}]"""
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    total_pages = len(reader.pages)
    pages_to_process = min(total_pages, MAX_PAGES)
    pages = []

    for i in range(pages_to_process):
        text = reader.pages[i].extract_text() or ""
        text = text[:MAX_TEXT_PER_PAGE].strip()
        if text:
            pages.append({"page_number": i + 1, "text": text})

    reader.close()
    del reader
    return pages


def extract_pages_docx(file_path: str) -> list:
    """Extrae contenido de un DOCX. Agrupa párrafos en páginas lógicas."""
    from docx import Document

    doc = Document(file_path)
    pages = []
    current_text = ""
    page_num = 1

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        current_text += text + "\n"

        # Agrupar en bloques de ~MAX_TEXT_PER_PAGE chars como "páginas"
        if len(current_text) >= MAX_TEXT_PER_PAGE:
            pages.append({"page_number": page_num, "text": current_text.strip()})
            current_text = ""
            page_num += 1

            if page_num > MAX_PAGES:
                break

    # Último bloque
    if current_text.strip() and page_num <= MAX_PAGES:
        pages.append({"page_number": page_num, "text": current_text.strip()})

    # Extraer también texto de tablas
    for table in doc.tables:
        if page_num > MAX_PAGES:
            break
        table_text = ""
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_text += " | ".join(row_cells) + "\n"

        if table_text.strip():
            pages.append({"page_number": page_num, "text": table_text.strip()[:MAX_TEXT_PER_PAGE]})
            page_num += 1

    return pages


def extract_pages_xlsx(file_path: str) -> list:
    """Extrae contenido de un Excel. Cada hoja = una página."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    pages = []
    page_num = 1

    for sheet_name in wb.sheetnames:
        if page_num > MAX_PAGES:
            break

        ws = wb[sheet_name]
        rows_text = []

        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows_text.append(" | ".join(cells))

        if rows_text:
            text = f"[Hoja: {sheet_name}]\n" + "\n".join(rows_text)
            text = text[:MAX_TEXT_PER_PAGE]
            pages.append({"page_number": page_num, "text": text})
            page_num += 1

    wb.close()
    return pages


def extract_pages_txt(file_path: str) -> list:
    """Extrae contenido de un archivo de texto plano."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    pages = []
    page_num = 1

    # Dividir en bloques de MAX_TEXT_PER_PAGE
    for i in range(0, len(content), MAX_TEXT_PER_PAGE):
        if page_num > MAX_PAGES:
            break

        block = content[i:i + MAX_TEXT_PER_PAGE].strip()
        if block:
            pages.append({"page_number": page_num, "text": block})
            page_num += 1

    return pages


# --- Procesamiento principal ---

EXTRACTORS = {
    ".pdf": extract_pages_pdf,
    ".docx": extract_pages_docx,
    ".xlsx": extract_pages_xlsx,
    ".txt": extract_pages_txt,
}


def process_file(file_path: str, filename: str, doc_type: str, obra: str, fecha: str, uploaded_by: str = None):
    """Procesa un archivo y lo ingesta en Supabase"""

    openai_key = os.getenv("OPENAI_API_KEY")
    supabase_url = os.getenv("SUPABASE_URL")
    supabase_key = os.getenv("SUPABASE_KEY")

    if not all([openai_key, supabase_url, supabase_key]):
        return {"success": False, "error": "Faltan variables de entorno"}

    # Detectar formato
    ext = Path(filename).suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if not extractor:
        return {"success": False, "error": f"Formato no soportado: {ext}"}

    try:
        supabase = create_client(supabase_url, supabase_key)

        # Extraer páginas/secciones del archivo
        pages = extractor(file_path)

        total_chunks = 0

        for page_data in pages:
            try:
                text = page_data["text"]
                page_num = page_data["page_number"]

                if not text:
                    continue

                chunks = split_text(text)

                for chunk_idx, chunk_text in enumerate(chunks):
                    if not chunk_text:
                        continue

                    embedding = get_openai_embedding(chunk_text, openai_key)
                    encrypted_content = encrypt_content(chunk_text)

                    doc = {
                        "content": encrypted_content,
                        "embedding": embedding,
                        "source": filename,
                        "page_number": page_num,
                        "document_type": doc_type,
                        "obra": obra if obra else None,
                        "fecha": fecha if fecha else None,
                        "metadata": {"chunk_index": chunk_idx},
                        "uploaded_by": uploaded_by
                    }

                    supabase.table("documents").insert(doc).execute()
                    total_chunks += 1

                    del embedding
                    del doc

                del chunks
                gc.collect()

            except Exception as e:
                print(f"Error en seccion {page_num}: {e}", file=sys.stderr)
                continue

        gc.collect()

        return {
            "success": True,
            "chunks_processed": total_chunks,
            "pages_processed": len(pages),
        }

    except Exception as e:
        return {"success": False, "error": str(e)}


def main():
    if len(sys.argv) < 3:
        print(json.dumps({"success": False, "error": "Argumentos insuficientes"}))
        sys.exit(1)

    file_path = sys.argv[1]
    filename = sys.argv[2]
    doc_type = sys.argv[3] if len(sys.argv) > 3 else "documento"
    obra = sys.argv[4] if len(sys.argv) > 4 and sys.argv[4] else ""
    fecha = sys.argv[5] if len(sys.argv) > 5 and sys.argv[5] else ""
    uploaded_by = sys.argv[6] if len(sys.argv) > 6 and sys.argv[6] else None

    result = process_file(file_path, filename, doc_type, obra, fecha, uploaded_by)
    print(json.dumps(result))

    if not result["success"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
